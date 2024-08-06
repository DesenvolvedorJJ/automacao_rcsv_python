import pandas as pd
from openpyxl import Workbook
from openpyxl.chart import BarChart, Reference
from openpyxl.chart.label import DataLabelList
from openpyxl.styles import Color, PatternFill
from openpyxl.utils.dataframe import dataframe_to_rows

# Carregar o arquivo Excel e ler os dados
df = pd.read_excel('Produção por site/07.xlsx')

# Verificar se as colunas necessárias estão presentes
if 'Site' in df.columns and 'Quantidade total' in df.columns:
    # Criar um novo Workbook
    wb = Workbook()
    ws = wb.active
    ws.title = 'Dados'

    # Adicionar os dados do DataFrame na planilha
    for r in dataframe_to_rows(df, index=False, header=True):
        ws.append(r)

    # Criar o gráfico de colunas horizontais
    chart = BarChart()
    chart.type = "bar"
    chart.style = 10
    chart.title = "QUANTIADADE TOTAL DE IMPRESSÕES POR PRÉDIO"
    chart.y_axis.title = ' '
    chart.x_axis.title = ' '

    # Adicionar os dados ao gráfico
    data = Reference(ws, min_col=2, min_row=1, max_row=len(df) + 1, max_col=2)
    categories = Reference(ws, min_col=1, min_row=2, max_row=len(df) + 1)
    chart.add_data(data, titles_from_data=True)
    chart.set_categories(categories)

    # Configurar a cor das colunas para amarela
    for series in chart.series:
        series.graphicalProperties.solidFill = "FFC106"  # Cor amarela

    # Adicionar rótulos de dados à direita das colunas
    chart.dLbls = DataLabelList()
    chart.dLbls.showVal = True
    chart.dLbls.dLblPos = "r"  # Posicionar os rótulos à direita das colunas

    # Remover a legenda do gráfico
    chart.legend = None

    # Adicionar o gráfico à planilha
    ws.add_chart(chart, "F1")

    # Salvar o arquivo Excel com o gráfico
    wb.save('prod_por_predio_07.xlsx')
else:
    print("Erro.")
